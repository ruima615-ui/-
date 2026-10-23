import hashlib
import shutil
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from sqlalchemy.orm import Session
from sqlalchemy import delete

from app.core.config import COURSE_MATERIAL_DIR
from app.models.entities import CourseMaterial
from app.models.entities import CourseMaterialChunkTopic
from app.models.entities import Course
from app.models.entities import User
from app.repositories.course_material_repository import (
    add_course_material,
    delete_course_material as delete_course_material_record,
    get_course_material,
    list_course_materials,
)
from app.services.teacher_service import (
    require_teacher_course_access,
)


ALLOWED_MATERIAL_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
    ".pptx",
}
MAX_MATERIAL_SIZE_BYTES = 20 * 1024 * 1024


def _extract_plain_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("文本资料必须使用 UTF-8 编码") from error


def _extract_pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise RuntimeError("缺少 PDF 解析依赖 pypdf") from error

    reader = PdfReader(BytesIO(data))
    return "\n\n".join(
        page.extract_text() or ""
        for page in reader.pages
    )


def _extract_docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as error:
        raise RuntimeError("缺少 DOCX 解析依赖 python-docx") from error

    document = Document(BytesIO(data))
    lines = []
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if not value:
            continue
        style = getattr(getattr(paragraph, "style", None), "name", "") or ""
        if style.lower().startswith("heading") or style.startswith("标题"):
            lines.append(f"## {value}")
        else:
            lines.append(value)
    return "\n".join(lines)


def _extract_pptx_text(data: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as error:
        raise RuntimeError("缺少 PPTX 解析依赖 python-pptx") from error

    presentation = Presentation(BytesIO(data))
    lines = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        lines.append(f"## 第 {slide_number} 页")
        first_text = True
        for shape in slide.shapes:
            text = getattr(shape, "text", None)
            if text:
                if first_text:
                    lines.append(f"### {text.strip()}")
                    first_text = False
                else:
                    lines.append(text)

    return "\n".join(lines)


def extract_material_text(
    data: bytes,
    extension: str,
) -> str:
    """根据扩展名提取资料文本，输出供后续分块使用的统一文本。"""

    extractors = {
        ".txt": _extract_plain_text,
        ".md": _extract_plain_text,
        ".pdf": _extract_pdf_text,
        ".docx": _extract_docx_text,
        ".pptx": _extract_pptx_text,
    }
    extractor = extractors[extension]
    text = extractor(data).strip()

    if not text:
        raise ValueError("资料中没有提取到可用文本")

    return text


def _validate_upload(
    filename: str,
    data: bytes,
) -> tuple[str, str]:
    safe_filename = Path(filename).name

    if not safe_filename:
        raise ValueError("上传文件必须包含文件名")
    if safe_filename != filename:
        raise ValueError("上传文件名不能包含目录路径")
    if not data:
        raise ValueError("不能上传空文件")
    if len(data) > MAX_MATERIAL_SIZE_BYTES:
        raise ValueError("课程资料不能超过 20MB")

    extension = Path(safe_filename).suffix.lower()
    if extension not in ALLOWED_MATERIAL_EXTENSIONS:
        supported = ", ".join(
            sorted(ALLOWED_MATERIAL_EXTENSIONS)
        )
        raise ValueError(
            f"不支持的资料格式，允许格式：{supported}"
        )

    return safe_filename, extension


def _create_course_material(
    *,
    db: Session,
    teacher_id: str,
    course_id: str,
    filename: str,
    mime_type: str | None,
    data: bytes,
    title: str | None = None,
    topic: str | None = None,
    storage_root: Path | None = None,
    source_type: str = "teacher",
) -> CourseMaterial:
    """保存并解析一份已经完成身份授权的课程资料。"""

    if source_type not in {"teacher", "official"}:
        raise ValueError("资料来源类型无效")
    safe_filename, extension = _validate_upload(
        filename,
        data,
    )
    resolved_title = (
        title
        if title is not None
        else Path(safe_filename).stem
    ).strip()

    if not resolved_title:
        raise ValueError("课程资料标题不能为空")
    if len(resolved_title) > 200:
        raise ValueError("课程资料标题不能超过 200 个字符")
    # 文档级字段只记录教师给出的候选范围；每个 Chunk 的实际标签在索引时确定。
    resolved_topic = (topic or "通用").strip()
    if not resolved_topic:
        resolved_topic = "通用"
    if len(resolved_topic) > 100:
        raise ValueError("课程资料知识点不能超过 100 个字符")

    material_id = f"MAT_{uuid4().hex}"
    root = storage_root or COURSE_MATERIAL_DIR
    material_dir = root / material_id
    material_dir.mkdir(parents=True, exist_ok=False)

    source_path = material_dir / f"source{extension}"
    extracted_path = material_dir / "extracted.txt"
    source_path.write_bytes(data)

    status = "processing"
    parse_error = None
    extracted_text = ""

    try:
        extracted_text = extract_material_text(
            data=data,
            extension=extension,
        )
        extracted_path.write_text(
            extracted_text,
            encoding="utf-8",
        )
        status = "ready_for_indexing"
    except (ValueError, RuntimeError, OSError) as error:
        status = "parse_failed"
        parse_error = str(error)

    material = CourseMaterial(
        id=material_id,
        course_id=course_id,
        uploaded_by=teacher_id,
        source_type=source_type,
        topic=resolved_topic,
        title=resolved_title,
        original_filename=safe_filename,
        file_extension=extension,
        mime_type=mime_type,
        size_bytes=len(data),
        checksum_sha256=hashlib.sha256(data).hexdigest(),
        source_path=str(source_path.relative_to(root)),
        extracted_text_path=(
            str(extracted_path.relative_to(root))
            if status == "ready_for_indexing"
            else None
        ),
        extracted_text_length=len(extracted_text),
        status=status,
        parse_error=parse_error,
    )

    try:
        add_course_material(db, material)
        db.commit()
        db.refresh(material)
        return material
    except Exception:
        db.rollback()
        raise


def create_teacher_course_material(
    *,
    db: Session,
    teacher_id: str,
    course_id: str,
    filename: str,
    mime_type: str | None,
    data: bytes,
    title: str | None = None,
    topic: str | None = None,
    storage_root: Path | None = None,
) -> CourseMaterial:
    """教师只能为自己负责的课程上传 teacher 类型资料。"""

    require_teacher_course_access(
        db=db,
        teacher_id=teacher_id,
        course_id=course_id,
    )
    return _create_course_material(
        db=db,
        teacher_id=teacher_id,
        course_id=course_id,
        filename=filename,
        mime_type=mime_type,
        data=data,
        title=title,
        topic=topic,
        storage_root=storage_root,
        source_type="teacher",
    )


def _require_active_admin_and_course(
    *,
    db: Session,
    admin_id: str,
    course_id: str,
) -> Course:
    admin = db.get(User, admin_id)
    if (
        admin is None
        or admin.role != "admin"
        or admin.status != "active"
    ):
        raise PermissionError("只有可用的管理员可以上传官方资料")

    course = db.get(Course, course_id)
    if course is None or course.status != "active":
        raise ValueError("课程不存在或不可用")

    return course


def create_official_course_material(
    *,
    db: Session,
    admin_id: str,
    course_id: str,
    filename: str,
    mime_type: str | None,
    data: bytes,
    title: str | None = None,
    topic: str | None = None,
    storage_root: Path | None = None,
) -> CourseMaterial:
    """管理员为指定有效课程上传 official 类型资料。"""

    _require_active_admin_and_course(
        db=db,
        admin_id=admin_id,
        course_id=course_id,
    )
    return _create_course_material(
        db=db,
        teacher_id=admin_id,
        course_id=course_id,
        filename=filename,
        mime_type=mime_type,
        data=data,
        title=title,
        topic=topic,
        storage_root=storage_root,
        source_type="official",
    )


def get_teacher_course_materials(
    *,
    db: Session,
    teacher_id: str,
    course_id: str,
) -> list[CourseMaterial]:
    """列出教师有权管理的课程资料。"""

    require_teacher_course_access(
        db=db,
        teacher_id=teacher_id,
        course_id=course_id,
    )
    return list_course_materials(
        db=db,
        course_id=course_id,
    )


def get_admin_course_materials(
    *,
    db: Session,
    admin_id: str,
    course_id: str,
) -> list[CourseMaterial]:
    """管理员查看指定有效课程的全部资料。"""

    _require_active_admin_and_course(
        db=db,
        admin_id=admin_id,
        course_id=course_id,
    )
    return list_course_materials(
        db=db,
        course_id=course_id,
    )


def _require_material_owner(
    *,
    db: Session,
    user_id: str,
    material_id: str,
    expected_source_type: str,
) -> CourseMaterial:
    material = get_course_material(db, material_id)
    if material is None:
        raise ValueError("课程资料不存在")

    if expected_source_type == "teacher":
        require_teacher_course_access(
            db=db,
            teacher_id=user_id,
            course_id=material.course_id,
        )
    else:
        _require_active_admin_and_course(
            db=db,
            admin_id=user_id,
            course_id=material.course_id,
        )

    if (
        material.source_type != expected_source_type
        or material.uploaded_by != user_id
    ):
        raise PermissionError("只能管理自己上传的同类型课程资料")

    return material


def update_owned_course_material_title(
    *,
    db: Session,
    user_id: str,
    material_id: str,
    title: str,
    expected_source_type: str,
) -> CourseMaterial:
    """教师或管理员修改自己上传的资料标题。"""

    material = _require_material_owner(
        db=db,
        user_id=user_id,
        material_id=material_id,
        expected_source_type=expected_source_type,
    )
    normalized_title = title.strip()
    if not normalized_title:
        raise ValueError("课程资料标题不能为空")

    try:
        material.title = normalized_title
        db.commit()
        db.refresh(material)
        return material
    except Exception:
        db.rollback()
        raise


def delete_owned_course_material(
    *,
    db: Session,
    user_id: str,
    material_id: str,
    expected_source_type: str,
    storage_root: Path | None = None,
    vector_store=None,
) -> None:
    """删除本人上传的资料记录和对应本地文件。"""

    material = _require_material_owner(
        db=db,
        user_id=user_id,
        material_id=material_id,
        expected_source_type=expected_source_type,
    )
    root = (storage_root or COURSE_MATERIAL_DIR).resolve()
    material_dir = (
        root / Path(material.source_path).parent
    ).resolve()

    if root != material_dir and root not in material_dir.parents:
        raise ValueError("课程资料存储路径无效")

    try:
        # 延迟导入，避免资料服务和索引服务形成模块循环。
        from app.services.material_index_service import (
            remove_course_material_from_index,
        )

        remove_course_material_from_index(
            db=db,
            material_id=material.id,
            vector_store=vector_store,
        )
        db.execute(delete(CourseMaterialChunkTopic).where(
            CourseMaterialChunkTopic.material_id == material.id))
        delete_course_material_record(db, material)
        db.commit()
    except Exception:
        db.rollback()
        raise

    if material_dir.is_dir():
        shutil.rmtree(material_dir)
