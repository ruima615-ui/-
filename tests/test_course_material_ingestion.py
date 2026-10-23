import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.exam_dependencies import get_current_user
from app.core.database import Base
from app.core.database import get_db
from app.main import app
from app.models.entities import Course
from app.models.entities import TeacherCourseAssignment
from app.models.entities import User
from app.services.course_material_service import (
    create_official_course_material,
    create_teacher_course_material,
    delete_owned_course_material,
    extract_material_text,
    get_teacher_course_materials,
    update_owned_course_material_title,
)


def build_minimal_text_pdf(text: str) -> bytes:
    """生成只用于解析器测试的最小英文文本 PDF。"""

    stream = (
        "BT /F1 12 Tf 72 720 Td "
        f"({text}) Tj ET"
    ).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R "
            b"/MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> "
            b"/Contents 4 0 R >>"
        ),
        (
            f"<< /Length {len(stream)} >>\nstream\n".encode("ascii")
            + stream
            + b"\nendstream"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]

    for number, body in enumerate(objects, start=1):
        offsets.append(len(result))
        result.extend(
            f"{number} 0 obj\n".encode("ascii")
        )
        result.extend(body)
        result.extend(b"\nendobj\n")

    xref_offset = len(result)
    result.extend(
        f"xref\n0 {len(objects) + 1}\n".encode("ascii")
    )
    result.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        result.extend(
            f"{offset:010d} 00000 n \n".encode("ascii")
        )
    result.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} "
            "/Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(result)


class CourseMaterialIngestionTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine)
        self.db = self.Session()
        self.temp_dir = tempfile.TemporaryDirectory()
        self.storage_root = Path(self.temp_dir.name)

        self.teacher = User(
            id="TEACHER_A",
            name="教师A",
            role="teacher",
            status="active",
        )
        self.other_teacher = User(
            id="TEACHER_B",
            name="教师B",
            role="teacher",
            status="active",
        )
        self.admin = User(
            id="ADMIN_A",
            name="管理员A",
            role="admin",
            status="active",
        )
        course_a = Course(
            id="COURSE_A",
            code="COURSE_A",
            name="课程A",
            status="active",
        )
        course_b = Course(
            id="COURSE_B",
            code="COURSE_B",
            name="课程B",
            status="active",
        )
        self.db.add_all([
            self.teacher,
            self.other_teacher,
            self.admin,
            course_a,
            course_b,
        ])
        self.db.flush()
        self.db.add_all([
            TeacherCourseAssignment(
                teacher_id="TEACHER_A",
                course_id="COURSE_A",
                assigned_by="ADMIN_A",
            ),
            TeacherCourseAssignment(
                teacher_id="TEACHER_B",
                course_id="COURSE_B",
                assigned_by="ADMIN_A",
            ),
        ])
        self.db.commit()

    def tearDown(self):
        app.dependency_overrides.clear()
        self.db.close()
        self.engine.dispose()
        self.temp_dir.cleanup()

    def test_markdown_upload_is_saved_and_ready_for_indexing(self):
        material = create_teacher_course_material(
            db=self.db,
            teacher_id="TEACHER_A",
            course_id="COURSE_A",
            filename="python_list.md",
            mime_type="text/markdown",
            data="# Python列表\n列表索引从0开始。".encode("utf-8"),
            storage_root=self.storage_root,
        )

        self.assertEqual(material.status, "ready_for_indexing")
        self.assertEqual(material.source_type, "teacher")
        self.assertEqual(material.course_id, "COURSE_A")
        self.assertGreater(material.extracted_text_length, 0)
        self.assertEqual(len(material.checksum_sha256), 64)
        self.assertTrue(
            (self.storage_root / material.source_path).is_file()
        )
        extracted_path = (
            self.storage_root / material.extracted_text_path
        )
        self.assertIn(
            "列表索引从0开始",
            extracted_path.read_text(encoding="utf-8"),
        )

    def test_teacher_cannot_upload_to_unassigned_course(self):
        with self.assertRaisesRegex(
            PermissionError,
            "无权管理这门课程",
        ):
            create_teacher_course_material(
                db=self.db,
                teacher_id="TEACHER_A",
                course_id="COURSE_B",
                filename="forbidden.md",
                mime_type="text/markdown",
                data=b"forbidden",
                storage_root=self.storage_root,
            )

        self.assertEqual(
            list(self.storage_root.iterdir()),
            [],
        )

    def test_unsupported_extension_is_rejected(self):
        with self.assertRaisesRegex(
            ValueError,
            "不支持的资料格式",
        ):
            create_teacher_course_material(
                db=self.db,
                teacher_id="TEACHER_A",
                course_id="COURSE_A",
                filename="program.exe",
                mime_type="application/octet-stream",
                data=b"not allowed",
                storage_root=self.storage_root,
            )

    def test_material_list_is_limited_to_authorized_course(self):
        create_teacher_course_material(
            db=self.db,
            teacher_id="TEACHER_A",
            course_id="COURSE_A",
            filename="a.md",
            mime_type="text/markdown",
            data="课程A资料".encode("utf-8"),
            storage_root=self.storage_root,
        )
        create_teacher_course_material(
            db=self.db,
            teacher_id="TEACHER_B",
            course_id="COURSE_B",
            filename="b.md",
            mime_type="text/markdown",
            data="课程B资料".encode("utf-8"),
            storage_root=self.storage_root,
        )

        materials = get_teacher_course_materials(
            db=self.db,
            teacher_id="TEACHER_A",
            course_id="COURSE_A",
        )

        self.assertEqual(len(materials), 1)
        self.assertEqual(materials[0].original_filename, "a.md")

    def test_teacher_upload_api_uses_authenticated_teacher(self):
        def override_get_db():
            yield self.db

        def override_current_user():
            return self.teacher

        app.dependency_overrides[get_db] = override_get_db
        app.dependency_overrides[
            get_current_user
        ] = override_current_user

        with (
            patch(
                "app.services.course_material_service."
                "COURSE_MATERIAL_DIR",
                self.storage_root,
            ),
            patch(
                "app.api.course_materials."
                "index_course_material_in_background",
            ) as index_mock,
            TestClient(app) as client,
        ):
            response = client.post(
                "/api/teacher/course-materials/courses/COURSE_A",
                data={
                    "title": "列表课程讲义",
                    "topic": "Python列表",
                },
                files={
                    "file": (
                        "list.md",
                        "# 列表\n索引从0开始。".encode("utf-8"),
                        "text/markdown",
                    )
                },
            )

        self.assertEqual(response.status_code, 201)
        payload = response.json()
        self.assertEqual(payload["course_id"], "COURSE_A")
        self.assertEqual(payload["uploaded_by"], "TEACHER_A")
        self.assertEqual(payload["source_type"], "teacher")
        self.assertEqual(payload["topic"], "Python列表")
        self.assertEqual(payload["title"], "列表课程讲义")
        self.assertEqual(payload["status"], "ready_for_indexing")
        index_mock.assert_called_once_with(payload["material_id"])

    def test_pdf_docx_and_pptx_text_extractors(self):
        from docx import Document
        from pptx import Presentation

        pdf_text = extract_material_text(
            build_minimal_text_pdf(
                "Python list index starts at zero."
            ),
            ".pdf",
        )

        document = Document()
        document.add_paragraph("DOCX中的Python列表资料")
        docx_buffer = BytesIO()
        document.save(docx_buffer)
        docx_text = extract_material_text(
            docx_buffer.getvalue(),
            ".docx",
        )

        presentation = Presentation()
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )
        slide.shapes.title.text = "PPTX列表索引"
        slide.placeholders[1].text = "索引从0开始"
        pptx_buffer = BytesIO()
        presentation.save(pptx_buffer)
        pptx_text = extract_material_text(
            pptx_buffer.getvalue(),
            ".pptx",
        )

        self.assertIn("starts at zero", pdf_text)
        self.assertIn("DOCX中的Python列表资料", docx_text)
        self.assertIn("PPTX列表索引", pptx_text)
        self.assertIn("索引从0开始", pptx_text)

    def test_admin_upload_api_creates_official_source(self):
        def override_get_db():
            yield self.db

        def override_current_user():
            return self.admin

        app.dependency_overrides[get_db] = override_get_db
        app.dependency_overrides[
            get_current_user
        ] = override_current_user

        with (
            patch(
                "app.services.course_material_service."
                "COURSE_MATERIAL_DIR",
                self.storage_root,
            ),
            patch(
                "app.api.course_materials."
                "index_course_material_in_background",
            ) as index_mock,
            TestClient(app) as client,
        ):
            response = client.post(
                "/api/admin/course-materials/courses/COURSE_A",
                files={
                    "file": (
                        "official.md",
                        "官方课程资料".encode("utf-8"),
                        "text/markdown",
                    )
                },
            )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(
            response.json()["source_type"],
            "official",
        )
        index_mock.assert_called_once_with(
            response.json()["material_id"]
        )

    def test_teacher_cannot_use_official_upload_api(self):
        def override_get_db():
            yield self.db

        def override_current_user():
            return self.teacher

        app.dependency_overrides[get_db] = override_get_db
        app.dependency_overrides[
            get_current_user
        ] = override_current_user

        with (
            patch(
                "app.services.course_material_service."
                "COURSE_MATERIAL_DIR",
                self.storage_root,
            ),
            TestClient(app) as client,
        ):
            response = client.post(
                "/api/admin/course-materials/courses/COURSE_A",
                files={
                    "file": (
                        "fake_official.md",
                        "不能标记为官方".encode("utf-8"),
                        "text/markdown",
                    )
                },
            )

        self.assertEqual(response.status_code, 403)

    def test_official_source_is_derived_from_admin_identity(self):
        material = create_official_course_material(
            db=self.db,
            admin_id="ADMIN_A",
            course_id="COURSE_A",
            filename="official.md",
            mime_type="text/markdown",
            data="官方课程资料".encode("utf-8"),
            storage_root=self.storage_root,
        )

        self.assertEqual(material.source_type, "official")
        self.assertEqual(material.uploaded_by, "ADMIN_A")

    def test_teacher_cannot_update_official_material(self):
        material = create_official_course_material(
            db=self.db,
            admin_id="ADMIN_A",
            course_id="COURSE_A",
            filename="official.md",
            mime_type="text/markdown",
            data="官方课程资料".encode("utf-8"),
            storage_root=self.storage_root,
        )

        with self.assertRaisesRegex(
            PermissionError,
            "只能管理自己上传",
        ):
            update_owned_course_material_title(
                db=self.db,
                user_id="TEACHER_A",
                material_id=material.id,
                title="不能修改",
                expected_source_type="teacher",
            )

    def test_teacher_can_update_and_delete_own_material(self):
        material = create_teacher_course_material(
            db=self.db,
            teacher_id="TEACHER_A",
            course_id="COURSE_A",
            filename="owned.md",
            mime_type="text/markdown",
            data="教师自己的资料".encode("utf-8"),
            storage_root=self.storage_root,
        )
        material_dir = (
            self.storage_root / Path(material.source_path).parent
        )

        updated = update_owned_course_material_title(
            db=self.db,
            user_id="TEACHER_A",
            material_id=material.id,
            title="更新后的标题",
            expected_source_type="teacher",
        )
        self.assertEqual(updated.title, "更新后的标题")

        delete_owned_course_material(
            db=self.db,
            user_id="TEACHER_A",
            material_id=material.id,
            expected_source_type="teacher",
            storage_root=self.storage_root,
        )

        self.assertIsNone(
            self.db.get(type(material), material.id)
        )
        self.assertFalse(material_dir.exists())


if __name__ == "__main__":
    unittest.main()
